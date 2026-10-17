import base64
import io
import os
import re
from datetime import datetime

import streamlit as st
from PIL import Image
from docx import Document
from docx.shared import Inches
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch

try:
    import fitz  # PyMuPDF: renders PDF pages for OCR
except ImportError:
    fitz = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

APP_NAME = "PaperScan AI"
st.set_page_config(page_title=APP_NAME, page_icon="📜", layout="wide")

st.markdown("""
<style>
    .stApp { background: linear-gradient(180deg, #f7fbff 0%, #ffffff 48%, #f4fbf7 100%); }
    .hero { padding: 1.35rem 1.5rem; border-radius: 22px; color: #123047;
            background: linear-gradient(115deg,#e4f3ff,#e8fff2); border: 1px solid #d6e8ef; }
    .hero h1 { margin: 0; font-size: 2.35rem; }
    .hero p { margin: .4rem 0 0 0; font-size: 1.05rem; }
    .small-note { color: #526777; font-size: .9rem; }
    div[data-testid="stMetric"] { background: white; padding: 12px; border-radius: 14px; border: 1px solid #e7edf1; }
</style>
""", unsafe_allow_html=True)

if "ocr_text" not in st.session_state:
    st.session_state.ocr_text = ""
if "chat" not in st.session_state:
    st.session_state.chat = []

def get_api_key():
    try:
        return st.secrets.get("OPENAI_API_KEY", "") or os.getenv("OPENAI_API_KEY", "")
    except Exception:
        return os.getenv("OPENAI_API_KEY", "")

def get_client():
    key = get_api_key()
    if not key or OpenAI is None:
        return None
    try:
        return OpenAI(api_key=key)
    except Exception:
        return None

def image_to_data_url(image: Image.Image) -> str:
    buf = io.BytesIO()
    image.convert("RGB").save(buf, format="JPEG", quality=90)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("utf-8")

def pdf_pages_to_images(pdf_bytes, max_pages=15, dpi=180):
    if fitz is None:
        raise RuntimeError("PyMuPDF is not installed. Add pymupdf to requirements.txt.")
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    images = []
    for page in doc[:max_pages]:
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
        images.append(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))
    return images, len(doc)

def ocr_with_openai(image: Image.Image, instructions: str) -> str:
    client = get_client()
    if client is None:
        raise RuntimeError("OpenAI API key/client is not configured.")
    model = st.session_state.get("vision_model", "gpt-4o-mini")
    response = client.responses.create(
        model=model,
        instructions=(
            "You are a careful OCR engine. Transcribe all visible text faithfully. "
            "Preserve paragraph breaks, headings, lists, tables where practical, and punctuation. "
            "For handwriting, infer only when reasonably clear; mark uncertain words as [unclear]. "
            "Do not summarize, correct grammar, or add commentary. " + instructions
        ),
        input=[{"role": "user", "content": [
            {"type": "input_text", "text": "Transcribe the text in this document image."},
            {"type": "input_image", "image_url": image_to_data_url(image)}
        ]}],
    )
    return response.output_text.strip()

def ocr_with_tesseract(image: Image.Image) -> str:
    if pytesseract is None:
        raise RuntimeError("Tesseract OCR Python package is missing.")
    return pytesseract.image_to_string(image).strip()

def extract_pdf_text(pdf_bytes):
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()

def make_docx(text, title="PaperScan AI document"):
    buf = io.BytesIO()
    doc = Document()
    doc.add_heading(title, 0)
    for block in re.split(r"\n\s*\n", text.strip()):
        if block.strip():
            p = doc.add_paragraph()
            for line_i, line in enumerate(block.splitlines()):
                if line_i:
                    p.add_run().add_break()
                p.add_run(line)
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def make_pdf(text, title="PaperScan AI document"):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=0.7*inch, leftMargin=0.7*inch,
                            topMargin=0.7*inch, bottomMargin=0.7*inch)
    styles = getSampleStyleSheet()
    story = [Paragraph(title.replace("&", "&amp;"), styles["Title"]), Spacer(1, 14)]
    for block in re.split(r"\n\s*\n", text.strip()):
        if block.strip():
            safe = (block.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    .replace("\n", "<br/>"))
            story.extend([Paragraph(safe, styles["BodyText"]), Spacer(1, 9)])
    doc.build(story)
    buf.seek(0)
    return buf.getvalue()

def ask_ai(prompt, text, model="gpt-4o-mini"):
    client = get_client()
    if client is None:
        raise RuntimeError("Add OPENAI_API_KEY to Streamlit Secrets or environment variables to use AI tools.")
    response = client.responses.create(
        model=model,
        instructions="You are PaperScan AI, a precise document editor and writing assistant. Follow the user's requested output format. Never invent facts or silently change names, dates, amounts, or technical details.",
        input=f"{prompt}\n\nDOCUMENT / CONTEXT:\n{text}",
    )
    return response.output_text.strip()

st.markdown("""
<div class="hero">
  <h1>📜 PaperScan AI</h1>
  <p>Turn handwritten notes, printed pages, and scanned documents into editable text and polished documents.</p>
</div>
""", unsafe_allow_html=True)
st.write("")
col1, col2, col3 = st.columns(3)
col1.metric("OCR workspace", "Image + PDF")
col2.metric("Export formats", "TXT · PDF · DOCX")
col3.metric("Writing tools", "Polish · Email · Chat")

with st.sidebar:
    st.header("⚙️ Settings")
    st.session_state.vision_model = st.selectbox("Vision/OCR model", ["gpt-4o-mini", "gpt-4.1-mini", "gpt-4o"], index=0)
    st.caption("AI OCR and writing tools require an OpenAI API key.")
    if get_api_key():
        st.success("OpenAI API key detected")
    else:
        st.info("No API key detected. Searchable PDFs may still yield embedded text; local Tesseract OCR works if installed.")
    st.markdown("---")
    st.caption("Designed for careful transcription. Always proofread handwriting, numbers, and names.")

tabs = st.tabs(["🔎 Scan & OCR", "✨ Polish with AI", "✉️ Email Formatter", "🔄 File Converter", "💬 Assistant"])

with tabs[0]:
    st.subheader("Scan handwritten or printed pages")
    uploaded = st.file_uploader("Upload an image or PDF", type=["png", "jpg", "jpeg", "webp", "tif", "tiff", "bmp", "pdf"], key="scan_upload")
    col_a, col_b = st.columns([1, 1])
    with col_a:
        method = st.radio("OCR method", ["AI OCR (best for handwriting)", "Local Tesseract OCR", "Extract embedded PDF text"], index=0)
    with col_b:
        extra = st.text_input("Optional context (e.g., English, Urdu, engineering notes)", "")
    if uploaded:
        file_bytes = uploaded.getvalue()
        is_pdf = uploaded.name.lower().endswith(".pdf")
        if is_pdf:
            st.info(f"PDF selected: {uploaded.name}")
            if fitz is not None:
                try:
                    imgs, total_pages = pdf_pages_to_images(file_bytes)
                    st.caption(f"{total_pages} page(s) found. OCR processes up to 15 pages per run.")
                    if imgs:
                        st.image(imgs[0], caption="Preview — first page", use_container_width=True)
                except Exception as e:
                    imgs = []
                    st.warning(f"Could not render PDF preview: {e}")
            else:
                imgs = []
        else:
            try:
                image = Image.open(io.BytesIO(file_bytes)).convert("RGB")
                st.image(image, caption=uploaded.name, use_container_width=True)
                imgs = [image]
            except Exception as e:
                imgs = []
                st.error(f"Could not open image: {e}")
        if st.button("🚀 Extract text", type="primary", use_container_width=True):
            try:
                if method == "Extract embedded PDF text":
                    if not is_pdf:
                        st.warning("Choose a PDF to extract embedded text.")
                    else:
                        result = extract_pdf_text(file_bytes)
                        if not result:
                            st.warning("No embedded text found. This may be a scanned PDF; select AI OCR or Tesseract instead.")
                        else:
                            st.session_state.ocr_text = result
                elif method == "AI OCR (best for handwriting)":
                    if get_client() is None:
                        st.error("AI OCR requires OPENAI_API_KEY. Add it in Streamlit Secrets, or use local Tesseract.")
                    else:
                        if is_pdf and not imgs:
                            imgs, _ = pdf_pages_to_images(file_bytes)
                        if not imgs:
                            st.error("No pages available to process.")
                        else:
                            output = []
                            progress = st.progress(0)
                            for idx, img in enumerate(imgs):
                                output.append(f"--- Page {idx + 1} ---\n" + ocr_with_openai(img, extra))
                                progress.progress((idx + 1) / len(imgs))
                            st.session_state.ocr_text = "\n\n".join(output).strip()
                else:
                    if is_pdf and not imgs:
                        imgs, _ = pdf_pages_to_images(file_bytes)
                    if not imgs:
                        st.error("No image pages available for local OCR.")
                    else:
                        output = [f"--- Page {idx + 1} ---\n{ocr_with_tesseract(img)}" for idx, img in enumerate(imgs)]
                        st.session_state.ocr_text = "\n\n".join(output).strip()
            except Exception as e:
                st.error(f"OCR failed: {e}")
    st.markdown("#### Extracted text")
    st.session_state.ocr_text = st.text_area("Review and edit OCR output before exporting", value=st.session_state.ocr_text, height=280, key="ocr_editor")
    text_out = st.session_state.ocr_text
    if text_out.strip():
        c1, c2, c3 = st.columns(3)
        c1.download_button("⬇️ Download TXT", data=text_out.encode("utf-8"), file_name="paperscan_text.txt", mime="text/plain", use_container_width=True)
        c2.download_button("⬇️ Download DOCX", data=make_docx(text_out), file_name="paperscan_document.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
        c3.download_button("⬇️ Download PDF", data=make_pdf(text_out), file_name="paperscan_document.pdf", mime="application/pdf", use_container_width=True)

with tabs[1]:
    st.subheader("Polish OCR text without losing its meaning")
    polish_source = st.text_area("Text to polish", value=st.session_state.ocr_text, height=220, key="polish_source")
    goal = st.selectbox("Choose an improvement", [
        "Fix OCR spacing and line breaks only",
        "Correct grammar and spelling",
        "Rewrite professionally while preserving meaning",
        "Create a clear summary",
        "Convert into structured notes with headings",
        "Translate into English",
        "Translate into Urdu",
    ])
    if st.button("✨ Improve content", type="primary"):
        try:
            result = ask_ai(f"Task: {goal}. Preserve all factual details, names, dates, numbers, and technical meaning. If uncertain, flag it rather than guessing.", polish_source, st.session_state.vision_model)
            st.session_state.polished = result
        except Exception as e:
            st.error(str(e))
    polished = st.text_area("Improved result", value=st.session_state.get("polished", ""), height=250, key="polished_editor")
    if polished.strip():
        if st.button("Use improved text in OCR workspace"):
            st.session_state.ocr_text = polished
            st.success("Improved text copied to the OCR workspace. Switch to Scan & OCR to export it.")
        st.download_button("Download polished TXT", polished.encode("utf-8"), "paperscan_polished.txt", "text/plain")

with tabs[2]:
    st.subheader("Professional email formatter")
    email_goal = st.selectbox("Email purpose", ["Formal request", "Follow-up", "Meeting invitation", "Thank-you email", "Complaint / issue report", "Application / cover email", "General professional email"])
    recipient = st.text_input("Recipient / organization (optional)")
    subject = st.text_input("Subject or topic")
    tone = st.selectbox("Tone", ["Professional and concise", "Warm and respectful", "Formal and detailed", "Friendly"])
    email_notes = st.text_area("Key points to include", height=170, placeholder="Write rough notes, voice transcription, or copied OCR text here…")
    if st.button("✉️ Format email", type="primary"):
        try:
            prompt = f"Create a ready-to-send email. Purpose: {email_goal}. Recipient: {recipient or 'not specified'}. Topic: {subject or 'derive a concise subject'}. Tone: {tone}. Include a subject line and email body. Do not invent commitments or details."
            st.session_state.formatted_email = ask_ai(prompt, email_notes, st.session_state.vision_model)
        except Exception as e:
            st.error(str(e))
    formatted_email = st.text_area("Formatted email", value=st.session_state.get("formatted_email", ""), height=260, key="email_result")
    if formatted_email.strip():
        st.download_button("Download email as TXT", formatted_email.encode("utf-8"), "formatted_email.txt", "text/plain")

with tabs[3]:
    st.subheader("Convert documents")
    st.caption("TXT and DOCX can be converted to PDF. PDFs with selectable text can be exported to TXT or DOCX.")
    converter_file = st.file_uploader("Choose a TXT, DOCX, or PDF file", type=["txt", "docx", "pdf"], key="converter_upload")
    if converter_file:
        raw = converter_file.getvalue()
        name = converter_file.name.lower()
        extracted = ""
        try:
            if name.endswith(".txt"):
                extracted = raw.decode("utf-8", errors="replace")
            elif name.endswith(".docx"):
                doc = Document(io.BytesIO(raw))
                extracted = "\n\n".join(p.text for p in doc.paragraphs).strip()
                for table in doc.tables:
                    extracted += "\n\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            elif name.endswith(".pdf"):
                extracted = extract_pdf_text(raw)
                if not extracted:
                    st.warning("This PDF appears to be scanned or image-only. Use Scan & OCR first.")
            st.text_area("Extracted content preview", value=extracted, height=220, key="convert_preview")
            if extracted.strip():
                left, middle, right = st.columns(3)
                left.download_button("TXT", extracted.encode("utf-8"), "converted.txt", "text/plain", use_container_width=True)
                middle.download_button("DOCX", make_docx(extracted, converter_file.name), "converted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                right.download_button("PDF", make_pdf(extracted, converter_file.name), "converted.pdf", "application/pdf", use_container_width=True)
        except Exception as e:
            st.error(f"Conversion failed: {e}")

with tabs[4]:
    st.subheader("PaperScan AI assistant")
    st.caption("Ask for help understanding a document, improving wording, extracting action items, or preparing a letter.")
    for item in st.session_state.chat:
        with st.chat_message(item["role"]):
            st.markdown(item["content"])
    user_msg = st.chat_input("Ask PaperScan AI…")
    if user_msg:
        st.session_state.chat.append({"role": "user", "content": user_msg})
        with st.chat_message("user"):
            st.markdown(user_msg)
        context = st.session_state.ocr_text or "(No scanned text in workspace yet.)"
        try:
            answer = ask_ai("Answer the user's question about the document or writing task. If context is insufficient, say what is missing. Be accurate and helpful.\nUSER QUESTION: " + user_msg, context, st.session_state.vision_model)
        except Exception as e:
            answer = f"AI assistant is unavailable: {e}"
        st.session_state.chat.append({"role": "assistant", "content": answer})
        with st.chat_message("assistant"):
            st.markdown(answer)

st.markdown("---")
st.markdown("<p class='small-note'>📜 PaperScan AI · Review OCR results before using them in official, legal, academic, or financial documents. Handwriting and low-quality scans may contain recognition errors.</p>", unsafe_allow_html=True)
